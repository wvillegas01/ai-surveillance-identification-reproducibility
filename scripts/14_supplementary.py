# -*- coding: utf-8 -*-
"""
Etapa 20l — Material suplementario (Supplementary Tables 1-8) en un libro
Excel autonomo: hoja de contenidos con los pies + una hoja por tabla, cada
una con su propio pie (Frontiers no maqueta el suplemento).

Salida: 15_secciones_manuscrito/Frontiers_LaTeX_Templates/Supplementary_Material/
        Supplementary_Material.xlsx

Uso:  python scripts/14_supplementary.py
"""
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import common as rc

DEST = rc.SUPP
TITULO = ("Identification Limits in Estimating the Effects of Algorithmic Surveillance "
          "on Digital Freedoms")
O = lambda f: pd.read_csv(os.path.join(rc.OUT, f))
OUTC = dict(rc.INDICADORES)
REGLA = {"year_surv_pos": "Positive surveillance tag (primary, Eq. 1)",
         "year_surv_neg": "Smart-city records excluded",
         "year_min": "Earliest record of any kind",
         "year_2nd": "Second distinct year",
         "year_p25": "25th percentile of qualifying years"}
COVS = {"sin covariables": "None", "dos covariables": "Two (electoral democracy, rule of law)",
        "cinco covariables": "Five"}
TRANS = {"log": "Logarithms", "niveles": "Levels", "-": "Not applicable"}
MUESTRA = {"completa": "Complete comparison group (77)",
           "sin marcados": "Comparison countries with records excluded (59)",
           "sin China": "China excluded"}
COVNOM = {"v2polyarchy": "Electoral democracy (v2x_polyarchy)",
          "v2x_polyarchy": "Electoral democracy (v2x_polyarchy)",
          "v2rule": "Rule of law (v2x_rule)", "v2x_rule": "Rule of law (v2x_rule)",
          "v2stfisccap": "Fiscal capacity (v2stfisccap)",
          "e_gdppc": "GDP per capita, log (e_gdppc)", "e_pop": "Population, log (e_pop)",
          "v2libdem": "Liberal democracy (v2x_libdem)",
          "v2x_libdem": "Liberal democracy (v2x_libdem)",
          "v2svstterr": "Territorial control (v2svstterr)",
          "v2freexp_altinf": "Freedom of expression (v2x_freexp_altinf)",
          "v2x_freexp_altinf": "Freedom of expression (v2x_freexp_altinf)"}
REGION = {1: "Eastern Europe and Central Asia", 2: "Latin America and the Caribbean",
          3: "Middle East and North Africa", 4: "Sub-Saharan Africa",
          5: "Western Europe and North America", 6: "Asia and Pacific"}
TECH = {"safe_city": "Safe city", "facial": "Facial recognition",
        "smart_policing": "Smart policing", "sm_monitoring": "Social media monitoring"}
TIPO = {"webpage": "Web page", "blogPost": "Blog post", "journalArticle": "Journal article",
        "newspaperArticle": "Newspaper article", "magazineArticle": "Magazine article",
        "report": "Report", "conferencePaper": "Conference paper", "book": "Book",
        "bookSection": "Book section", "document": "Document"}
NOTA = ("All outcomes are latent scores from the Digital Society Project v8; higher values "
        "indicate greater digital freedom (less repression). Estimation sample: 158 countries, "
        "2000-2022; 81 dated documented adopters in eight cohorts and 77 countries without "
        "documented adoption. Primary specification: Callaway-Sant'Anna outcome regression, "
        "base period g-1, five covariates averaged over 2000-2011 (GDP per capita and "
        "population in logarithms), differences 0.3.0 (Python 3.12).")


def st1():
    e = O("event_study_todos.csv")
    e = e[e.escenario == "principal"].copy()
    d = pd.DataFrame({
        "Outcome": e.variable.map(OUTC), "Variable": e.variable,
        "Event time (e = t - g)": e.rel.astype(int),
        "Estimate": e.att, "SE (analytical)": e.se, "SE (bootstrap)": e.se_bootstrap,
        "95% CI lower": e.lo, "95% CI upper": e.hi,
        "Simultaneous band lower": e.lo_sim, "Simultaneous band upper": e.hi_sim,
        "Pointwise CI excludes zero": e.excluye_cero_puntual.map({True: "Yes", False: "No"}),
        "Simultaneous band excludes zero": e.excluye_cero_simult.map({True: "Yes", False: "No"}),
        "Contributing cohorts": e.n_cohortes.astype(int),
        "Contributing treated countries": e.n_paises.astype(int)})
    cap = ("Supplementary Table 1. Complete event-study estimates for the eight outcomes under "
           "the primary specification (event times -18 to 10). Pre-treatment coefficients use "
           "consecutive differences; post-treatment coefficients use base period g-1. "
           "Simultaneous bands are Bonferroni-adjusted within the pre-treatment (18) and "
           "post-treatment (11) families. Bootstrap SE: 499 country-level replications "
           "stratified by treatment status.")
    return d, cap


def st2():
    b = O("balance_por_cohorte.csv")
    d = pd.DataFrame({"Cohort": b.cohorte, "Covariate": b.covariable.map(COVNOM).fillna(b.covariable),
                      "In adjustment model": b.en_el_modelo.map({True: "Yes", False: "No (external diagnostic)"}),
                      "SMD before weighting": b.smd_antes, "SMD after weighting": b.smd_despues})
    cap = ("Supplementary Table 2. Cohort-specific covariate balance between treated and "
           "comparison countries before and after propensity-score weighting (logit with five "
           "covariates, income and population in logarithms; odds weights for comparison "
           "countries). Covariates are means over 2000-2011. SMD: standardized mean difference "
           "with the unweighted pooled standard deviation in the denominator.")
    return d, cap


def fecha_iso(raw):
    """Normaliza la fecha registrada en la biblioteca a ISO 8601 sin inventar
    precision: devuelve (fecha, precision)."""
    import re
    from datetime import datetime
    s = str(raw).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return "-".join(m.groups()), "day"
    m = re.match(r"^(\d{4})/(\d{2})/(\d{2})$", s)
    if m:
        return "-".join(m.groups()), "day"
    m = re.match(r"^\w{3}, (\d{2})/(\d{2})/(\d{4})", s)      # "Sun, 11/03/2013 - 09:07"
    if m:
        return "%s-%s-%s" % (m.group(3), m.group(1), m.group(2)), "day"
    for fmt in ("%B %d, %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d"), "day"
        except ValueError:
            pass
    for fmt in ("%B %Y", "%b %Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m"), "month (day not recorded)"
        except ValueError:
            pass
    m = re.match(r"^(\d{4})$", s)
    if m:
        return m.group(1), "year (month and day not recorded)"
    raise ValueError("fecha no reconocida: %r" % s)


def st3():
    ev = O("evidencias_por_pais.csv")
    reg = O("reglas_por_pais.csv")
    ev = ev[ev.fija_la_fecha == True].merge(
        reg[["country_text_id"] + rc.REGLAS], on="country_text_id", how="left")
    ev = ev[ev.rol == "tratado_datado"]
    cod = O("codificacion_estado_operativo.csv")
    ev = ev.drop(columns=["estado_operativo", "evidencia_textual",
                          "anio_despliegue_segun_fuente"]).merge(
        cod[["country_text_id", "key", "estado_operativo", "evidencia_textual",
             "anio_despliegue_segun_fuente", "base_codificacion"]],
        on=["country_text_id", "key"], how="left")
    assert ev.estado_operativo.notna().all()
    zot = pd.read_csv(rc.ZOTERO)[["key", "date", "year"]]
    ev = ev.merge(zot, on="key", how="left")
    fp = ev.date.apply(fecha_iso)
    ev["fecha_iso"], ev["precision"] = fp.str[0], fp.str[1]
    assert (ev.fecha_iso.str[:4].astype(int) == ev.year.astype(int)).all()
    assert (ev.year.astype(int) == ev.anio_publicacion.astype(int)).all()
    cod2 = O("codificacion_estado_operativo.csv")[["country_text_id", "key", "anio_mencionado",
                                                   "referencia_anio", "nota"]]
    ev = ev.merge(cod2, on=["country_text_id", "key"], how="left")
    ev["ident"] = ev.url
    ev.loc[ev.url.isna() & (ev.key == "UBF2EAGF"), "ident"] = "https://doi.org/10.1109/ICISCE.2017.103"
    d = pd.DataFrame({
        "Country": ev.country_name, "ISO3": ev.country_text_id,
        "First documented year (primary rule)": ev.anio_tratamiento.astype(int),
        "Year: smart-city records excluded": ev.year_surv_neg,
        "Year: earliest record of any kind": ev.year_min,
        "Year: second distinct year": ev.year_2nd,
        "Year: 25th percentile": ev.year_p25,
        "Zotero record key": ev.key,
        "Publication date (ISO 8601)": ev.fecha_iso,
        "Date precision": ev.precision,
        "Publication date as recorded in the library": ev.date,
        "Document type": ev.tipo.map(TIPO).fillna(ev.tipo), "Title": ev.titulo, "URL or identifier": ev.ident,
        "Technologies tagged": ev.tecnologias.fillna("").apply(
            lambda s: "; ".join(TECH.get(t, t) for t in s.split(";") if t)),
        "Countries covered by the record": ev.n_paises_item,
        "Operational status (implemented / contract / announcement / project / retrospective)":
            ev.estado_operativo.fillna(""),
        "Supporting evidence from the source": ev.evidencia_textual.fillna(""),
        "Year or period mentioned in the source": ev.anio_mencionado.fillna(""),
        "What the year refers to": ev.referencia_anio.fillna(""),
        "Basis of classification": ev.base_codificacion,
        "Note": ev.nota.fillna("")})
    cap = ("Supplementary Table 3. Country-level treatment evidence. For each of the 81 dated "
           "documented adopters, the bibliographic record(s) from the open library supporting "
           "the Global AI Surveillance Index that fix the first documented year under Equation 1 "
           "(several rows when more than one qualifying record shares that year), together with "
           "the year assigned under each alternative dating rule. The first documented year "
           "records publication, not deployment; the operational-status columns classify "
           "whether each record documents an implemented system, a contract or procurement, an "
           "announcement, a proposed project, or retrospective reporting. Coding protocol "
           "(fixed before reading the sources): Implemented = the source reports a system "
           "installed, in use, deployed or launched, including running pilots; Contract = a "
           "signed agreement, sale, purchase or award without evidence of operation; "
           "Announcement = plans, negotiations, memoranda, legal authorisation or announced "
           "future use; Project = an approved or prepared project without evidence of a "
           "contract or operation; Retrospective = a later source describing an earlier "
           "deployment that can be dated; Not established = the source is not accessible, does "
           "not mention the country, or does not document a state deployment, in which case "
           "nothing is inferred. Sources were read on 2026-10-12 in full text (original or "
           "Internet Archive) or, when not accessible, through the abstract stored in the "
           "library. Quotations are brief extracts. A year is recorded only when the source "
           "states it, together with what it refers to (operation, installation, pilot or "
           "test, planned operation, purchase, agreement, implementation work); only "
           "operation, installation and pilot years indicate deployment. Publication dates "
           "are normalised to ISO 8601 with their recorded precision; the publication year "
           "used for dating equals the year of the recorded date in every row.")
    return d, cap


def st4():
    m = rc.maestra()
    u = m[m.rol == "tratado_sin_datar"].copy()
    d = pd.DataFrame({
        "Country": u.country_name, "ISO3": u.country_text_id,
        "Region": u.region.map(REGION),
        "Technologies in the index": u.aigs_dose.astype(int),
        "Safe city": u.tech_safecity.map({1: "Yes", 0: "No"}),
        "Facial recognition": u.tech_facial.map({1: "Yes", 0: "No"}),
        "Smart policing": u.tech_smartpolicing.map({1: "Yes", 0: "No"}),
        "Social media monitoring": u.tech_socialmedia.map({1: "Yes", 0: "No"}),
        "Chinese supplier": u.vendor_china.map({1: "Yes", 0: "No"}),
        "U.S. supplier": u.vendor_us.map({1: "Yes", 0: "No"}),
        "Dated records of any kind": u.n_fechadas.fillna(0).astype(int),
        "Earliest record of any kind (year)": u.year_min})
    s = O("sin_datar_descriptivo.csv")
    cap = ("Supplementary Table 4. Documented adopters without a dated qualifying record (15 "
           "countries), excluded from the staggered-adoption estimation. Compared with the 81 "
           "dated adopters they are documented with fewer technologies (mean %.2f vs %.2f), "
           "are less often associated with facial recognition (%.0f%% vs %.0f%%) and U.S. "
           "suppliers (%.0f%% vs %.0f%%), and have smaller populations and lower incomes."
           % tuple(s.set_index("variable").loc[["aigs_dose", "aigs_dose", "tech_facial",
                                               "tech_facial", "vendor_us", "vendor_us"]]
                   .pipe(lambda x: [x.iloc[0].media_sin_datar, x.iloc[1].media_datados,
                                    100 * x.iloc[2].media_sin_datar, 100 * x.iloc[3].media_datados,
                                    100 * x.iloc[4].media_sin_datar, 100 * x.iloc[5].media_datados])))
    return d, cap


def st5():
    mv = O("multiverso_v2.csv")
    d = pd.DataFrame({"Dating rule": mv.regla_datacion.map(REGLA),
                      "Covariates": mv.covariables.map(COVS),
                      "Income and population": mv.transformacion.map(TRANS),
                      "Comparison sample": mv.muestra.map(MUESTRA),
                      "Panel end": mv.fin_panel, "Estimate": mv.ATT, "SE": mv.SE,
                      "95% CI lower": mv.IC_inf, "95% CI upper": mv.IC_sup, "p-value": mv.p,
                      "Cohorts with estimates": mv.cohortes_vivas,
                      "Treated countries": mv.n_tratados,
                      "Comparison countries": mv.n_comparacion})
    cap = ("Supplementary Table 5. Specification analysis for internet filtering: 75 "
           "configurations (5 dating rules x 3 covariate sets x 2 transformations x 3 comparison "
           "samples, with the transformation not applicable without covariates). All "
           "configurations use outcome regression and end in 2022. Estimate: cohort-size-"
           "weighted post-treatment summary with analytical SE.")
    return d, cap


def st6():
    filas = []
    r = O("resumen_escenarios.csv")
    esc = {"principal": "Primary specification",
           "sin controles marcados": "Comparison countries with records excluded",
           "covariables en niveles": "Income and population in levels",
           "ventana de covariables por cohorte": "Covariates in cohort-specific windows (2000 to g-1)"}
    for _i, x in r.iterrows():
        filas.append(["Specification", esc[x.escenario], OUTC[x.variable], x.ATT, x.SE, x.IC_inf,
                      x.IC_sup, x.p, x.n_tratados, x.cohortes_vivas])
    for _i, x in O("desplazamiento_fechas.csv").iterrows():
        if x.desplazamiento == 0:
            continue
        filas.append(["Treatment-date shift", "%d year%s earlier" % (x.desplazamiento,
                      "s" if x.desplazamiento > 1 else ""), OUTC[x.variable], x.ATT, x.SE,
                      x.IC_inf, x.IC_sup, x.p, x.n_tratados, x.cohortes_vivas])
    for _i, x in O("por_tecnologia.csv").iterrows():
        filas.append(["Technology-specific dating", TECH[x.familia], OUTC[x.variable], x.ATT,
                      x.SE, x.IC_inf, x.IC_sup, x.p, x.n_tratados, x.cohortes_vivas])
    for _i, x in O("sensibilidad_operativa.csv").iterrows():
        if x.muestra.startswith("Operational"):
            filas.append(["Operational evidence", "Treated restricted to countries whose dating "
                          "record documents an operating system", OUTC[x.variable], x.ATT, x.SE,
                          x.IC_inf, x.IC_sup, x.p, x.n_tratados, x.cohortes])
    for _i, x in O("soporte_restringido.csv").iterrows():
        filas.append(["Common support", "Within observed propensity-score overlap (bootstrap SE)",
                      OUTC[x.variable], x.ATT, x.SE, x.IC_inf, x.IC_sup, x.p, x.n_tratados,
                      x.n_cohortes])
    d = pd.DataFrame(filas, columns=["Analysis", "Variant", "Outcome", "Estimate", "SE",
                                     "95% CI lower", "95% CI upper", "p-value",
                                     "Treated countries", "Cohorts"])
    cap = ("Supplementary Table 6. Sensitivity analyses for the eight outcomes: alternative "
           "specifications, treatment dates shifted one to three years earlier (covariate window "
           "shifted accordingly), treated countries restricted to those whose dating record "
           "documents an operating system (Supplementary Table 3), technology-specific dating (each family dated by its own first "
           "qualifying record; other documented adopters excluded), and the subset within the "
           "observed propensity-score overlap. Estimates are cohort-size-weighted post-treatment "
           "summaries; SE analytical unless stated.")
    return d, cap


def st7():
    out = []
    for panel, lab in [("corregida", "Revised specification"), ("enviada", "Earlier specification")]:
        for est, el in [("reg", "Outcome regression"), ("dr", "Doubly robust")]:
            c = pd.read_csv(os.path.join(rc.OUT, "software", "celdas_comparadas_%s_%s.csv" % (panel, est)))
            c.insert(0, "Estimator", el); c.insert(0, "Specification", lab)
            if "manual" not in c:
                c["manual"] = np.nan
            out.append(c)
    c = pd.concat(out, ignore_index=True)
    d = pd.DataFrame({"Specification": c.Specification, "Estimator": c.Estimator,
                      "Cohort": c.cohort, "Period": c.time,
                      "differences 0.1.2 (Python 3.8)": c.v012,
                      "differences 0.3.0 (Python 3.12)": c.v030,
                      "csdid 0.4.2 (Python 3.12)": c.csdid,
                      "Direct implementation (outcome regression)": c.manual})
    cap = ("Supplementary Table 7. Cell-by-cell comparison of group-time estimates for internet "
           "filtering across implementations, with identical data, base periods, comparison "
           "group and covariates. Revised specification: positive-tag dating, covariates "
           "2000-2011 in logarithms. Earlier specification: earliest-record dating, covariates "
           "2000-2015 in levels. Empty cells are cohort-period estimates that the implementation "
           "returns as missing; for the doubly robust estimator they occur in cohorts where a "
           "treated country has a fitted propensity score numerically equal to one.")
    return d, cap


def st8():
    p = O("propension_por_pais.csv")
    mx = p[p.tratado == 0].groupby(["especificacion", "cohorte"]).ps.max().rename("max_c")
    p = p.join(mx, on=["especificacion", "cohorte"])
    nom = {"niveles": "Covariates in levels", "logaritmos": "Logarithms (primary)",
           "log + cuadraticos": "Logarithms with quadratic terms"}
    d = pd.DataFrame({"Logit specification": p.especificacion.map(nom), "Cohort": p.cohorte,
                      "Country": p.country_name, "ISO3": p.country_text_id,
                      "Group": p.tratado.map({1: "Treated", 0: "Comparison"}),
                      "Propensity score": p.ps,
                      "Treated above comparison maximum": np.where(
                          p.tratado == 1, np.where(p.ps > p.max_c, "Yes", "No"), ""),
                      "Outside [0.1, 0.9]": np.where(
                          p.tratado == 1, np.where((p.ps < .1) | (p.ps > .9), "Yes", "No"), "")})
    cap = ("Supplementary Table 8. Propensity scores by country and cohort under three logit "
           "specifications (five covariates in levels, in logarithms, and in logarithms with "
           "quadratic terms for income and population), fitted separately for each cohort on "
           "its treated countries and the 77 comparison countries.")
    return d, cap


def escribir(wb, nombre, d, cap):
    ws = wb.create_sheet(nombre)
    f = Font(name="Arial", size=10); fb = Font(name="Arial", size=10, bold=True)
    ws["A1"] = cap; ws["A1"].font = fb
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ncol = len(d.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncol, 6))
    ws.row_dimensions[1].height = 15 * (1 + len(cap) // 140)
    ws["A2"] = NOTA if nombre not in ("Supplementary Table 3", "Supplementary Table 4") else \
        "Source: open bibliographic library of the Global AI Surveillance Index " \
        "(zotero.org/groups/2347403), retrieved 2026-08-26; tags as assigned in the library."
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncol, 6))
    ws.row_dimensions[2].height = 40
    gris = PatternFill("solid", start_color="D9DEE7")
    borde = Border(bottom=Side(style="thin", color="7F8794"))
    for j, c in enumerate(d.columns, 1):
        cell = ws.cell(row=4, column=j, value=c)
        cell.font = fb; cell.fill = gris; cell.border = borde
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 45
    for i, fila in enumerate(d.itertuples(index=False), 5):
        for j, v in enumerate(fila, 1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = f
            if isinstance(v, float):
                col = d.columns[j - 1]
                cell.number_format = "0.0000" if not col.startswith("Year") else "0"
    for j, c in enumerate(d.columns, 1):
        largo = max([len(str(c)) // 2] + [len(str(x)) for x in d.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 60)
    ws.freeze_panes = "A5"
    return ws


def main():
    os.makedirs(DEST, exist_ok=True)
    tablas = [st1(), st2(), st3(), st4(), st5(), st6(), st7(), st8()]
    wb = Workbook(); ws = wb.active; ws.title = "Contents"
    fb = Font(name="Arial", size=12, bold=True); f = Font(name="Arial", size=10)
    ws["A1"] = "Supplementary Material"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = TITULO; ws["A2"].font = fb
    ws["A4"] = NOTA; ws["A4"].font = f
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:C4"); ws.row_dimensions[4].height = 75
    ws["A6"] = "Sheet"; ws["B6"] = "Rows"; ws["C6"] = "Caption"
    for c in ("A6", "B6", "C6"):
        ws[c].font = Font(name="Arial", size=10, bold=True)
    for k, (d, cap) in enumerate(tablas, 1):
        nombre = "Supplementary Table %d" % k
        escribir(wb, nombre, d, cap)
        r = 6 + k
        ws.cell(row=r, column=1, value=nombre).font = f
        ws.cell(row=r, column=2, value=len(d)).font = f
        c = ws.cell(row=r, column=3, value=cap); c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (1 + len(cap) // 110)
        print("  %-22s %5d filas x %2d columnas" % (nombre, len(d), len(d.columns)))
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 110
    salida = os.path.join(DEST, "Supplementary_Material.xlsx")
    wb.save(salida)
    print("-> %s" % salida)


if __name__ == "__main__":
    main()
